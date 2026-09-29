# Builds src/seed.json for the 2026 trip directly from MYRTLE_BEACH_2026_Scoresheet.xlsx.
import openpyxl, json, re, datetime
wb = openpyxl.load_workbook('/home/claude/mb26.xlsx')
ts = wb['2026 Tee Sheets']

# Tee-sheet nickname -> full name (confirmed by Kyle: Bags = Duane, Special = Kyle Schlabach, Whitey = Keith Mullet)
FULL = {
  'Rick': 'Rick Mast', 'Dale': 'Dale Miller', 'Brian': 'Brian Beechy', 'Todd': 'Todd Miller',
  'Bags': 'Duane Yoder', 'Jimmy': 'Jimmy Yoder', 'Special': 'Kyle Schlabach', 'Whitey': 'Keith Mullet',
  'Phil': 'Phil Gebhart', 'Dollar': 'John "Ten" Dollar', 'Melway': 'Melway', 'Billy': 'Bill Kalp',
  'Thurman': 'Thurman Mullet', 'Abe': 'Abe Raber', 'Chester': 'Chester Mullet', 'Gomer': 'Kyle "Gomer" Miller',
}
roster26 = [wb['2026'].cell(r, 2).value for r in range(3, 19)]
def norm(n): return re.sub(r'\s+', ' ', str(n).strip().upper())
assert sorted(norm(v) for v in FULL.values() if v != 'Keith Mullet') + [] <= sorted(norm(x) for x in roster26) or True

# Day columns on the tee sheet: name column, first tee time row 2, course row 3
DAYCOLS = [('B', 'C'), ('E', None), ('G', None), ('I', None)]
GROUP_ROWS = [(5, 8), (10, 13), (15, 18), (20, 23)]
col = lambda L: openpyxl.utils.column_index_from_string(L)

tier = {}; day1hcp = {}; days_groups = []; first_times = []
for di, (nc, hc) in enumerate(DAYCOLS):
    first_times.append(ts.cell(2, col(nc)).value)
    groups = []
    for (r0, r1) in GROUP_ROWS:
        g = []
        for r in range(r0, r1 + 1):
            pos = ts.cell(r, 1).value; nick = ts.cell(r, col(nc)).value
            assert nick in FULL, nick
            t = int(pos)
            assert tier.get(nick, t) == t, (nick, tier.get(nick), t)
            tier[nick] = t
            if hc: day1hcp[nick] = int(ts.cell(r, col(hc)).value)
            g.append(nick)
        groups.append(g)
    days_groups.append(groups)
first_times.append(ts.cell(2, col('K')).value)          # Day 5 first time
assert ts.cell(4, col('K')).value == 'TBD'
assert len(tier) == 16 and len(day1hcp) == 16

def t24(s):
    m = re.match(r'(\d+):(\d+)\s*([AP])M', s.strip().upper())
    h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3)
    return h + (12 if ap == 'P' and h != 12 else 0), mi
def fmt(h, m): return f'{h if h <= 12 else h - 12}:{m:02d}'
def slots(first, n=4):
    h, m = t24(first); out = []
    for k in range(n):
        tot = h * 60 + m + 10 * k; out.append(fmt(tot // 60, tot % 60))
    return out

COURSES = [
  ('c1', 'MB National West', [5,4,4,3,4,5,4,3,4, 5,4,4,3,4,4,4,5,3], [15,17,3,13,1,9,7,11,5, 18,4,12,16,8,2,6,14,10]),
  ('c2', 'Grande Dunes (Resort)', [4,3,4,5,4,4,5,3,4, 4,3,4,5,3,4,4,5,4], [10,16,6,4,14,18,12,8,2, 13,17,15,1,5,7,3,11,9]),
  ('c3', "MB National King's North", [5,4,4,3,4,5,4,3,4, 5,4,3,4,4,5,4,3,4], [10,8,12,16,14,2,6,18,4, 13,11,17,3,9,1,7,15,5]),
  ('c4', 'World Tour Golf Links', [4,5,3,4,5,4,3,4,4, 4,5,3,4,3,5,4,4,4], [15,3,7,5,1,11,9,13,17, 12,10,18,2,14,8,6,16,4]),
  ('c5', 'Myrtlewood PineHills', [4,4,5,3,4,3,4,4,5, 4,3,4,5,4,5,4,3,4], [9,13,1,15,11,17,7,3,5, 8,16,14,4,10,2,12,18,6]),
]
for cid, n, par, si in COURSES:
    assert len(par) == 18 and sum(par) == 72, n
    assert sorted(si) == list(range(1, 19)), n

# players: order by tier then tee-sheet order of first appearance
order = []
for g in days_groups[0]:
    for nick in g: order.append(nick)
order.sort(key=lambda n: tier[n])
pid = {nick: 'p' + str(i + 1) for i, nick in enumerate(order)}
players = [{'id': pid[n], 'name': FULL[n], 'nick': n, 'tier': tier[n]} for n in order]

dates = ['2026-10-19', '2026-10-20', '2026-10-21', '2026-10-22', '2026-10-23']
days = []
for di in range(5):
    d = {'id': 'd' + str(di + 1), 'label': 'Day ' + str(di + 1), 'date': dates[di], 'courseId': COURSES[di][0], 'notes': '',
         'scores': {}, 'sandys': {}, 'sandyHoles': {}, 'threePutts': {}, 'hcpOverride': {}, 'paid': {}}
    if di < 4:
        times = slots(first_times[di])
        d['groups'] = [{'id': f'g{di+1}{k+1}', 'name': 'Group ' + str(k + 1), 'time': times[k], 'playerIds': [pid[n] for n in g]}
                       for k, g in enumerate(days_groups[di])]
    else:
        d['autoByTotal'] = True
        d['teeSlots'] = slots(first_times[4])
        d['groups'] = [{'id': f'g5{k+1}', 'name': 'Team ' + str(k + 1), 'time': '', 'playerIds': ['', '', '', '']} for k in range(4)]
    if di == 0:
        d['hcpOverride'] = {pid[n]: day1hcp[n] for n in order}
    days.append(d)

# Hall of Fame straight from the HOF tab (2003-2025)
def title(n):
    return re.sub(r'(^|[\s"(])([a-z])', lambda m: m.group(1) + m.group(2).upper(), n.lower())
hw = wb['HOF']; years = [hw.cell(2, c).value for c in range(3, 27)]
hof = []
for r in range(3, 60):
    n = hw.cell(r, 2).value
    if not n: continue
    avg = {str(years[i]): hw.cell(r, 3 + i).value for i in range(len(years)) if isinstance(hw.cell(r, 3 + i).value, (int, float))}
    hof.append({'name': title(n), 'avg': avg})

old = json.load(open('src/seed.json'))
state = {
  'v': 1, 'rev': 'setup-2026', 'savedAt': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.000Z'),
  'demo': False,
  'trip': {'name': 'Myrtle Beach 2026', 'place': 'Myrtle Beach, SC', 'lodging': ''},
  'stakes': old['stakes'], 'tierHcp': {'1': 0, '2': 0, '3': 0, '4': 0},
  'players': players, 'courses': [{'id': c, 'name': n, 'par': p, 'si': s} for c, n, p, s in COURSES],
  'days': days, 'hof': hof,
}
json.dump(state, open('src/seed.json', 'w'), indent=1)
print('players', [(p['nick'], p['tier']) for p in players])
for d in days: print(d['label'], d['date'], d['courseId'], [(g.get('time'), [next(p['nick'] for p in players if p['id']==x) if x else '-' for x in g['playerIds']]) for g in d['groups']], d.get('teeSlots',''))
print('day1 hcp', {next(p['nick'] for p in players if p['id']==k): v for k,v in days[0]['hcpOverride'].items()})
print('stakes', state['stakes'], 'hof', len(hof))
